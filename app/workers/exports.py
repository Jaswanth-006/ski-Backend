"""Excel export builder (Phase 4-D). Pure function: day sheet → .xlsx bytes."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.schemas.day_sheet import DaySheetOut


def build_day_sheet_xlsx(sheet: DaySheetOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Day Sheet"

    ws["A1"] = f"Day Sheet — {sheet.business_date}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "CLOSED" if sheet.is_closed else "OPEN"

    headers = ["Delivery staff", "Loaded", "Sold", "Returned", "Cash", "UPI", "Total"]
    ws.append([])
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for row in sheet.rows:
        ws.append(
            [
                row.delivery_name,
                row.loaded,
                row.cylinders,
                row.returned,
                float(row.cash),
                float(row.upi),
                float(row.total),
            ]
        )

    ws.append(
        [
            "Total",
            sheet.stock.loaded,
            sheet.totals.cylinders,
            sheet.stock.returned,
            float(sheet.totals.cash),
            float(sheet.totals.upi),
            float(sheet.totals.total),
        ]
    )
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    ws.append([])
    ws.append([f"Warehouse stock — opening {sheet.stock.opening}, closing {sheet.stock.closing}"])
    ws.append([f"Cashier box — opening {sheet.cashier_opening}, closing {sheet.cashier_closing}"])
    ws.append(["Expenses", "", "", "", "", "", float(sheet.expenses_total)])
    ws.append(["Net (collection − expenses)", "", "", "", "", "", float(sheet.net)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    # Cash denomination breakdown (aggregated across all drivers).
    if sheet.denomination_totals:
        ws.append([])
        ws.append(["Cash denominations"])
        ws[ws.max_row][0].font = Font(bold=True)
        ws.append(["Note", "Count", "Amount"])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for d in sheet.denomination_totals:
            ws.append([d.note_value, d.note_count, d.note_value * d.note_count])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
